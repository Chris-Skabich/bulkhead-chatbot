import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from app.db import models

def generate_leads_excel(db: Session, client_urgency_threshold: float = 1.0) -> str:
    """
    Queries the database for all leads, sorts them by urgency, and highlights 
    critical leads that fall under the client's custom urgency threshold.
    """
    # Sort by timeline_months first (shortest/most urgent at the top)
    # Tie-breaker: largest linear feet job first
    leads = db.query(models.Lead).order_by(
        models.Lead.timeline_months.asc(), 
        models.Lead.linear_feet.desc()
    ).all()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Bulkhead Project Leads"
    ws.views.sheetView[0].showGridLines = True
    
    # Styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    
    # Define an Urgent Fill Color (Light Red) for critical rows
    urgent_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
    
    # Added 'Timeline (Raw)' and 'Timeline (AI Months)' 
    headers = ["ID", "Name", "Phone", "Project Type", "Linear Feet", "Timeline (Raw)", "Timeline (AI Months)", "Notes", "Date Created"]
    ws.append(headers)
    
    # Format header row
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    
    # Populate data rows
    for lead in leads:
        formatted_date = lead.created_at.strftime("%Y-%m-%d %H:%M") if lead.created_at else "N/A"
        
        # Protect against None values before doing math
        months = getattr(lead, 'timeline_months', 99.0)
        months = months if months is not None else 99.0
        
        row = [
            lead.id,
            lead.name,
            lead.phone,
            getattr(lead, 'project_type', 'N/A'),
            lead.linear_feet if lead.linear_feet is not None else "N/A",
            getattr(lead, 'timeline_raw', 'N/A'), # What the user actually typed
            months,                               # The AI converted number
            lead.notes if lead.notes else "",
            formatted_date
        ]
        ws.append(row)
        
        # HIGHLIGHT LOGIC: If this job is urgent, color the whole row red
        if months <= client_urgency_threshold:
            current_row = ws.max_row
            for col_idx in range(1, len(row) + 1):
                ws.cell(row=current_row, column=col_idx).fill = urgent_fill
    
    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    output_dir = "/tmp/reports"
    os.makedirs(output_dir, exist_ok=True)
    
    filename = f"leads_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    file_path = os.path.join(output_dir, filename)
    
    wb.save(file_path)
    return file_path